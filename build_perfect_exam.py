import os
import docx
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image, ImageDraw, ImageFont
import win32com.client
import pymupdf

FONT_PATH = "C:/Windows/Fonts/kaiu.ttf"

# High-resolution large handwriting box
# Box: W_BOX = 130, H = 130
# Zhuyin column: W_ZY = 54
# Total width: 184 px, height: 130 px
W_BOX = 130
W_ZY = 54
H = 130

f_zy = ImageFont.truetype(FONT_PATH, 36)
f_tone = ImageFont.truetype(FONT_PATH, 30)
f_hz = ImageFont.truetype(FONT_PATH, 82)

os.makedirs("exam_boxes_large", exist_ok=True)

def generate_large_box_img(char, zhuyin_str, mode, out_path):
    img = Image.new("RGB", (W_BOX + W_ZY, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([2, 2, W_BOX, H - 3], outline=(0, 0, 0), width=3)
    d.rectangle([W_BOX, 2, W_BOX + W_ZY, H - 3], outline=(0, 0, 0), width=3)
    
    tones = ['ˊ', 'ˇ', 'ˋ', '˙']
    tone = ''
    body = []
    for c in zhuyin_str:
        if c in tones:
            tone = c
        else:
            body.append(c)
            
    if mode == "char_blank":
        n = len(body)
        step = (H - 24) / max(n, 1)
        for i, c in enumerate(body):
            y = 10 + i * step
            d.text((W_BOX + 6, y), c, fill=(0, 0, 0), font=f_zy)
        if tone:
            if tone == '˙':
                d.text((W_BOX + 14, 2), tone, fill=(0, 0, 0), font=f_tone)
            else:
                tone_y = 10 + (n - 0.7) * step - 14
                d.text((W_BOX + 32, tone_y), tone, fill=(0, 0, 0), font=f_tone)
    elif mode == "zhuyin_blank":
        bbox = d.textbbox((0, 0), char, font=f_hz)
        cw, ch = bbox[2] - bbox[0], bbox[3] - bbox[1]
        x = (W_BOX - cw) / 2
        y = (H - ch) / 2 - 10
        d.text((x, y), char, fill=(0, 0, 0), font=f_hz)
        
    img.save(out_path)
    return out_path

items_p1 = [
    ("彷", "ㄈㄤˇ", "char_blank"),
    ("鷹", "ㄧㄥ", "zhuyin_blank"),
    ("衝", "ㄔㄨㄥ", "char_blank"),
    ("竟", "ㄐㄧㄥˋ", "zhuyin_blank"),
    ("訝", "ㄧㄚˋ", "char_blank"),
    ("恐", "ㄎㄨㄥˇ", "zhuyin_blank"),
    ("樹", "ㄕㄨˋ", "char_blank"),
    ("蕉", "ㄐㄧㄠ", "zhuyin_blank"),
    ("狀", "ㄓㄨㄤˋ", "char_blank"),
    ("隊", "ㄉㄨㄟˋ", "zhuyin_blank"),
    ("寧", "ㄋㄧㄥˊ", "char_blank"),
    ("控", "ㄎㄨㄥˋ", "zhuyin_blank"),
    ("制", "ㄓˋ", "char_blank"),
    ("碑", "ㄅㄟ", "zhuyin_blank"),
    ("巡", "ㄒㄩㄣˊ", "char_blank"),
    ("甜", "ㄊㄧㄢˊ", "zhuyin_blank"),
    ("航", "ㄏㄤˊ", "zhuyin_blank"),
    ("糖", "ㄊㄤˊ", "char_blank"),
    ("榮", "ㄖㄨㄥˊ", "zhuyin_blank"),
    ("衛", "ㄨㄟˋ", "char_blank"),
    ("灣", "ㄨㄢ", "zhuyin_blank"),
    ("嶼", "ㄩˇ", "char_blank"),
    ("翔", "ㄒㄧㄤˊ", "char_blank"),
    ("研", "ㄧㄢˊ", "zhuyin_blank"),
    ("造", "ㄗㄠˋ", "char_blank"),
    ("篇", "ㄆㄧㄢ", "zhuyin_blank"),
    ("盤", "ㄆㄢˊ", "char_blank"),
    ("耀", "ㄧㄠˋ", "zhuyin_blank"),
]

box_paths = {}
for i, (ch, zy, mode) in enumerate(items_p1):
    path = f"exam_boxes_large/box_{i}_{ch}_{mode}.png"
    generate_large_box_img(ch, zy, mode, path)
    box_paths[(ch, mode)] = path

doc = Document()
ns = nsdecls("w")

style_normal = doc.styles["Normal"]
style_normal.font.name = "標楷體"
rPr_normal = style_normal.element.get_or_add_rPr()
rFonts_normal = parse_xml(f'<w:rFonts {ns} w:ascii="標楷體" w:hAnsi="標楷體" w:eastAsia="標楷體" w:cs="標楷體" w:hint="eastAsia"/>')
rPr_normal.append(rFonts_normal)

sec = doc.sections[0]
# JIS B4: 364mm x 257mm (橫向)
sec.page_width = Mm(364)
sec.page_height = Mm(257)
sec.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
sec.left_margin = Mm(18)
sec.right_margin = Mm(18)
sec.top_margin = Mm(15)
sec.bottom_margin = Mm(15)

sectPr = sec._sectPr
sectPr.append(parse_xml(f'<w:textDirection {ns} w:val="tbRl"/>'))
docGrid = sectPr.find(qn('w:docGrid'))
if docGrid is not None:
    docGrid.set(qn('w:type'), 'default')
else:
    sectPr.append(parse_xml(f'<w:docGrid {ns} w:type="default"/>'))

def add_kai_run(p, text, size=15, bold=False):
    r = p.add_run(text)
    r.font.name = "標楷體"
    r.font.size = Pt(size)
    r.bold = bold
    rPr = r._element.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {ns} w:ascii="標楷體" w:hAnsi="標楷體" w:eastAsia="標楷體" w:cs="標楷體" w:hint="eastAsia"/>')
    rPr.append(rFonts)
    return r

def add_p(text="", size=15, bold=False, line_spacing=30, space_after=0):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:snapToGrid {ns} w:val="0"/>'))
    line_dxa = int(line_spacing * 20)
    pPr.append(parse_xml(f'<w:spacing {ns} w:line="{line_dxa}" w:lineRule="atLeast" w:after="{int(space_after * 20)}" w:before="0"/>'))
    if text:
        add_kai_run(p, text, size=size, bold=bold)
    return p

# Large box: width=46pt (across column), height=32pt (along column)
def add_box_run(p, ch, mode, width_pt=46, height_pt=32):
    path = box_paths[(ch, mode)]
    r = p.add_run()
    r.add_picture(path, width=Pt(width_pt), height=Pt(height_pt))
    return r

# ==================== PAGE 1 (第一面) ====================
add_p("臺中市梧棲區中正國小 115 學年度第一學期四年級國語科第一次定期成績評量試卷", size=15, line_spacing=26)
add_p("校長：＿＿＿    教務主任：＿＿＿    命題者：＿＿＿    審題者：四年級命題小組", size=15, line_spacing=26)
add_p("　　　　　　　　　四年　　班　　號姓名：＿＿＿＿＿＿　第一面（共四面）", size=15, line_spacing=26)

add_p("一、寫出國字和注音（共28分，每格1分）", size=15, bold=True, line_spacing=34)

# 行距設為 48pt，確保大格子（46pt寬）與相鄰欄位有完美充裕的安全間隙
p1_1 = add_p(line_spacing=48)
add_kai_run(p1_1, "　　臺中梧棲的【中正國小】即將舉辦校慶，四年級各班都在操場上集合。")

p1_2 = add_p(line_spacing=48)
add_kai_run(p1_2, "小明")
add_box_run(p1_2, "彷", "char_blank")
add_kai_run(p1_2, "彿是一隻")
add_box_run(p1_2, "鷹", "zhuyin_blank")
add_kai_run(p1_2, "，在跑道上全力")
add_box_run(p1_2, "衝", "char_blank")
add_kai_run(p1_2, "刺。沒想到，在交接棒時，他")
add_box_run(p1_2, "竟", "zhuyin_blank")
add_kai_run(p1_2, "然不小心失誤，把棒子掉在地上。全場觀眾發出驚")
add_box_run(p1_2, "訝", "char_blank")
add_kai_run(p1_2, "的聲音，小明也感到十分")
add_box_run(p1_2, "恐", "zhuyin_blank")
add_kai_run(p1_2, "懼與自責。")

p2_1 = add_p(line_spacing=48)
add_kai_run(p2_1, "　　比賽結束後，小明沮喪的坐在")
add_box_run(p2_1, "樹", "char_blank")
add_kai_run(p2_1, "下，默默吃著同學遞來的香")
add_box_run(p2_1, "蕉", "zhuyin_blank")
add_kai_run(p2_1, "。同學們見")
add_box_run(p2_1, "狀", "char_blank")
add_kai_run(p2_1, "，不但沒有責怪他，反而成群結")
add_box_run(p2_1, "隊", "zhuyin_blank")
add_kai_run(p2_1, "的走過來。班長輕輕拍著他的肩膀說：「別難過，我們是一個團體，要一起承擔。」")

p3_1 = add_p(line_spacing=48)
add_kai_run(p3_1, "　　小明聽了，原本")
add_box_run(p3_1, "寧", "char_blank")
add_kai_run(p3_1, "靜的心中感受到一絲溫暖。他努力")
add_box_run(p3_1, "控", "zhuyin_blank")
add_box_run(p3_1, "制", "char_blank")
add_kai_run(p3_1, "住自己的情緒，擦乾眼淚。這次的失敗，就像是刻在石")
add_box_run(p3_1, "碑", "zhuyin_blank")
add_kai_run(p3_1, "上的教訓，讓他學會了面對挫折。")

p4_1 = add_p(line_spacing=48)
add_kai_run(p4_1, "　　每天放學後，他都會留在操場上")
add_box_run(p4_1, "巡", "char_blank")
add_kai_run(p4_1, "視跑道，反覆練習接棒動作，直到汗水濕透了衣服。他相信，只要堅持下去，未來的日子一定會更加")
add_box_run(p4_1, "甜", "zhuyin_blank")
add_kai_run(p4_1, "蜜。")

p5_1 = add_p(line_spacing=48)
add_kai_run(p5_1, "　　在四年級這趟學習")
add_box_run(p5_1, "航", "zhuyin_blank")
add_kai_run(p5_1, "程中，我們不僅認識了臺灣甜蜜的")
add_box_run(p5_1, "糖", "char_blank")
add_kai_run(p5_1, "業文化與光")
add_box_run(p5_1, "榮", "zhuyin_blank")
add_kai_run(p5_1, "歷史，更體會了共同守")
add_box_run(p5_1, "衛", "char_blank")
add_kai_run(p5_1, "美麗海")
add_box_run(p5_1, "灣", "zhuyin_blank")
add_kai_run(p5_1, "與島")
add_box_run(p5_1, "嶼", "char_blank")
add_kai_run(p5_1, "的責任。")

p6_1 = add_p(line_spacing=48)
add_kai_run(p6_1, "　　我們渴望在天空翱")
add_box_run(p6_1, "翔", "char_blank")
add_kai_run(p6_1, "，追尋心中的夢想。只要用心觀察、努力")
add_box_run(p6_1, "研", "zhuyin_blank")
add_kai_run(p6_1, "究與改")
add_box_run(p6_1, "造", "char_blank")
add_kai_run(p6_1, "，每一")
add_box_run(p6_1, "篇", "zhuyin_blank")
add_kai_run(p6_1, "生活日記都能像銀")
add_box_run(p6_1, "盤", "char_blank")
add_kai_run(p6_1, "般閃")
add_box_run(p6_1, "耀", "zhuyin_blank")
add_kai_run(p6_1, "著希望的光芒。")

add_p("※後面還有試題，請翻面繼續作答。", size=15, bold=True, line_spacing=30)

doc.add_page_break()

# ==================== PAGE 2 (第二面) ====================
add_p("　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　第二面（共四面）", size=15, line_spacing=26)
add_p("二、改錯字：圈出錯誤的字，並將正確的字寫在【　】中（每字２分，共12分）", size=15, bold=True, line_spacing=28)

p_corr = add_p(line_spacing=28)
add_kai_run(p_corr, "　　中正國小的校園裡，有著許多美麗的風光。曾經有一次，小華因為考試成績不理想，心中十分自則。他獨自坐在操場旁邊，滿惱子都是負面的想法。這時，老師走了過來，溫柔的對他說：「面對挫拆並不可怕，重要的是你如何去面對它。」老師的話，彷服是一陣溫暖的微風，吹散了他心中的陰霾。小華決定不再距絕別人的幫助，他開始主動請教同學，並且學會了空制自己的脾氣。經過一段時間的努力，他的成績有了明顯的進步，臉上也重新煥發出自信的笑容。")

add_p("１【　】  ２【　】  ３【　】  ４【　】  ５【　】  ６【　】", size=15, line_spacing=28)

add_p("三、選擇題測驗（共20分，每題2分）", size=15, bold=True, line_spacing=28)

choices_p2 = [
    ("１（　）", "根據〈美麗島〉一課，詩中用「吐著唾沫」描寫海浪的手法效果是什麼？\n○１讓海浪動作生動活潑，像有生命一樣 ○２暗示海浪受到污染\n○３強調海浪拍打聲音巨大 ○４表示海浪退潮速度快。"),
    ("２（　）", "在〈請到我的家鄉來〉一課中，安安在信裡介紹了家鄉。如果小明要寫信介紹梧棲，信件最後應加上什麼才完整？\n○１景點與美食推薦 ○２祝福語、署名和日期\n○３歡迎來訪路線 ○４自己的聯絡方式。"),
    ("３（　）", "在〈鏡頭下的家鄉〉中，作者拍下許多人物照片，主要傳達什麼訊息？\n○１大家很喜歡拍照 ○２人物穿著很特別\n○３人們辛勤工作，展現生命力 ○４大家生活很富裕。"),
    ("４（　）", "關於〈飛行夢〉課文，下列哪一個選項的說明最正確？\n○１萊特兄弟從小接受嚴格科學訓練 ○２試飛過程中曾考慮放棄\n○３面對失敗依然堅持改良最後成功 ○４關閉了原本的腳踏車店。"),
    ("５（　）", "在〈月光下〉一課引用「小時不識月，呼作白玉盤」，下列哪句使用類似手法？\n○１妹妹臉頰紅紅的像熟透的蘋果 ○２星星好像在對我們眨眼睛\n○３他每天瘋狂練習像生病一樣 ○４安靜得連針掉地上都聽得到。")
]

for num, qtext in choices_p2:
    p_q = add_p(line_spacing=26)
    add_kai_run(p_q, num, bold=True)
    add_kai_run(p_q, qtext)

add_p("※後面還有試題，請翻面繼續作答。", size=15, bold=True, line_spacing=26)

doc.add_page_break()

# ==================== PAGE 3 (第三面) ====================
add_p("　　　　　　　　　四年　　班　　號姓名：＿＿＿＿＿＿　第三面（共四面）", size=15, line_spacing=26)

choices_p3 = [
    ("６（　）", "閱讀報紙時，想在最短時間知道報導核心重點，最應優先讀哪部分？\n○１新聞提要與目錄 ○２新聞照片及圖說\n○３大篇幅且醒目的標題 ○４報導日期與記者姓名。"),
    ("７（　）", "關於自然段與意義段的差別，下列敘述何者正確？\n○１自然段是換行退兩格段落，意義段合併相關自然段\n○２自然段重字數，意義段重句子長短\n○３一篇文章通常只有一個意義段 ○４兩者意思完全相同。"),
    ("８（　）", "「即使天氣不好，我們也照常舉行運動會」與哪一句意思最接近？\n○１因為天氣變差所以延期 ○２就算天氣變差還是照常舉行\n○３如果天氣不好就不舉行 ○４不但天氣差而且取消了。"),
    ("９（　）", "「成群結隊／形單影隻」的語詞關係，和下列哪一組相同？\n○１大開眼界／大飽口福 ○２花紅柳綠／風和日麗\n○３大排長龍／門可羅雀 ○４未雨綢繆／防患未然。"),
    ("10（　）", "當我們遇到挫折時，下列哪一種做法是比較適當的情緒調節方式？\n○１假裝一點都不在意 ○２立刻反省並嚴格責備自己\n○３找信任的家人或老師傾訴尋求幫助 ○４在網路上批評別人。")
]

for num, qtext in choices_p3:
    p_q = add_p(line_spacing=26)
    add_kai_run(p_q, num, bold=True)
    add_kai_run(p_q, qtext)

add_p("四、成語填空（共10分，每題2分）", size=15, bold=True, line_spacing=28)
add_p("請從下方成語選項中選出最適合填入句子的代號：\n○１多采多姿  ○２垂涎三尺  ○３家喻戶曉  ○４未雨綢繆  ○５大吃一驚  ○６大排長龍", size=15, line_spacing=26)

idioms = [
    "１ 媽媽煮的紅燒肉香氣四溢，讓人看了忍不住（　　）。",
    "２ 知名歌手來開演唱會，售票口總是（　　），一票難求。",
    "３ 颱風季節快到了，我們應該（　　），先準備好防災物品。",
    "４ 聽到他決定放棄出國留學的消息，大家都感到（　　）。",
    "５ 這次的校慶活動安排了各式各樣的表演，真是（　　）。"
]
for idm in idioms:
    add_p(idm, line_spacing=26)

add_p("※後面還有試題，請翻面繼續作答。", size=15, bold=True, line_spacing=26)

doc.add_page_break()

# ==================== PAGE 4 (第四面) ====================
add_p("　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　第四面（共四面）", size=15, line_spacing=26)
add_p("五、句型練習（共14分）", size=15, bold=True, line_spacing=28)
add_p("（一）照樣寫短語（每題3分，共6分）", size=15, line_spacing=26)
add_p("１（ 又遠又近 ）的（ 月亮 ）\n  答：（　　　　　　）的（　　　　　　）", size=15, line_spacing=28)
add_p("２（ 陽光 ）如（ 金粉 ）般（ 落下 ）\n  答：（　　　　　　）如（　　　　　　）般（　　　　　　）", size=15, line_spacing=28)

add_p("（二）造句（每題4分，共8分）", size=15, line_spacing=26)
add_p("１ 不只……而且…… ──\n  答：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿", size=15, line_spacing=28)
add_p("２ 即使……也…… ──\n  答：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿", size=15, line_spacing=28)

add_p("六、閱讀理解（共16分，每題4分）", size=15, bold=True, line_spacing=28)
add_p("請閱讀下列文章，並回答問題：", size=15, line_spacing=26)

p_read = add_p(line_spacing=26)
add_kai_run(p_read, "　　在廣大的火星表面，有一輛名為「毅力號」的探測車，正孤獨卻堅定的執行著任務。火星環境惡劣，溫差極大且常有沙塵暴。「毅力號」的主要任務是尋找古代生命遺跡並收集岩石樣本。然而在探索過程中，採樣管被一顆碎石卡住了。科學家們沒有放棄，經過冷靜分析與模擬，發送震動指令成功排除了碎石。「毅力號」再度恢復運作。這告訴我們，面對困難時，冷靜分析與堅持不懈是克服難關的關鍵。")

read_qs = [
    "１（　）「毅力號」在火星上的主要任務是什麼？\n○１尋找水源 ○２測量氣溫 ○３尋找古代生命遺跡並收集樣本 ○４建立通訊站。",
    "２（　）火星的環境有什麼特點？\n○１氣候溫和 ○２溫差極大且常有沙塵暴 ○３植物茂盛 ○４每天下大雨。",
    "３（　）採樣管被卡住時，科學家是如何解決問題的？\n○１發射救援車 ○２放棄採樣 ○３分析模擬並發送震動指令 ○４等待風吹走。",
    "４（　）這篇文章主要傳達了什麼道理？\n○１太空探索太危險 ○２面對困難應冷靜分析、堅持不懈 ○３機器一定會故障 ○４團隊合作不重要。"
]
for rq in read_qs:
    add_p(rq, line_spacing=26)

add_p("寫完後，記得再檢查喔！ 祝你有個好成績。 我預估我的國語可以考（      ）分", size=15, bold=True, line_spacing=26)

out_docx = "d:/opencode/115-/test_b4_perfect.docx"
doc.save(out_docx)
print("Saved test_b4_perfect.docx")
