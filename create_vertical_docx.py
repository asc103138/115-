import sys
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_vertical_exam(input_md, output_docx):
    doc = Document()
    
    # 設置版面 B4 橫式 (36.4 cm x 25.7 cm)
    section = doc.sections[0]
    section.page_width = Cm(36.4)
    section.page_height = Cm(25.7)
    section.orientation = WD_ORIENT.LANDSCAPE
    
    # 邊界設定
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    
    # 設定雙欄與直書
    sectPr = section._sectPr
    
    # 文字方向：直書由右至左
    text_dir = OxmlElement('w:textDirection')
    text_dir.set(qn('w:val'), 'tbRl')
    sectPr.insert(0, text_dir)
    
    # 雙欄與分隔線
    cols = sectPr.xpath('./w:cols')
    if cols:
        col_elem = cols[0]
    else:
        col_elem = OxmlElement('w:cols')
        sectPr.append(col_elem)
    col_elem.set(qn('w:num'), '2')
    col_elem.set(qn('w:space'), '720')  # 1.27 cm 欄間距
    col_elem.set(qn('w:sep'), '1')    # 分隔線
    
    # 全域字型：標楷體
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = '標楷體'
    normal_font.size = Pt(12)
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), '標楷體')
    
    # 段落格式
    pPr = normal_style._element.get_or_add_pPr()
    pPr.set(qn('w:snapToGrid'), '0')
    
    with open(input_md, 'r', encoding='utf-8') as f:
        content = f.read()
        
    lines = content.split('\n')
    for line in lines:
        if line.startswith('## 教師版標準解答'):
            # 遇到解答區就不放進學生考卷
            break
            
        # 簡單解析 markdown
        if line.startswith('#'):
            p = doc.add_paragraph(line.lstrip('#').strip())
            p.runs[0].font.size = Pt(16)
            p.runs[0].bold = True
        elif line.startswith('---'):
            # 處理分頁
            doc.add_page_break()
        else:
            p = doc.add_paragraph(line)
            if p.runs:
                p.runs[0].font.size = Pt(12)
            
    doc.save(output_docx)

if __name__ == '__main__':
    create_vertical_exam('期中評量初稿.md', '期中評量_排版版.docx')
