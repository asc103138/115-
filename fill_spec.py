import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client
import pymupdf

def fill_spec_table():
    src_docx = r"d:\opencode\115-\template_spec.docx"
    doc = Document(src_docx)
    ns = nsdecls("w")
    
    # Set default font to 標楷體
    style = doc.styles["Normal"]
    style.font.name = "標楷體"
    style.font.size = Pt(12)
    style._element.get_or_add_rPr().append(
        parse_xml(f'<w:rFonts {ns} w:ascii="標楷體" w:hAnsi="標楷體" w:eastAsia="標楷體" w:cs="標楷體" w:hint="eastAsia"/>')
    )
    
    # 1. Update Paragraphs
    # P0: Title
    # P2: 科目名稱：國語            命題教師：四年級命題小組            使用年級：四年級
    for p in doc.paragraphs:
        if "科目名稱" in p.text:
            p.text = "科目名稱：國語            命題教師：四年級命題小組            使用年級：四年級"
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "標楷體"
                r.font.size = Pt(12)
                r.bold = True
                r._element.get_or_add_rPr().append(
                    parse_xml(f'<w:rFonts {ns} w:ascii="標楷體" w:hAnsi="標楷體" w:eastAsia="標楷體" w:cs="標楷體" w:hint="eastAsia"/>')
                )
    
    def set_cell_text(cell, text, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER):
        cell.text = text
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = align
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "標楷體"
                r.font.size = Pt(size)
                r.bold = bold
                r._element.get_or_add_rPr().append(
                    parse_xml(f'<w:rFonts {ns} w:ascii="標楷體" w:hAnsi="標楷體" w:eastAsia="標楷體" w:cs="標楷體" w:hint="eastAsia"/>')
                )

    # 2. Table 1: 命題範圍之單元名稱、教學目標及教材比例分佈表
    t1 = doc.tables[0]
    t1_data = [
        ("第一單元：好山好水好故鄉", "認識美麗寶島、家鄉景物、書信互動與審題立意", "21", "47", "47"),
        ("　第 1 課 美麗島", "體會詩歌節奏韻律、動態描寫與寶島情感", "5", "11", "11"),
        ("　第 2 課 請到我的家鄉來", "閱讀書信內容、掌握信件結構與景點特色", "6", "13", "13"),
        ("　第 3 課 鏡頭下的家鄉", "透過攝影鏡頭觀察家鄉人物、描寫生活細節", "6", "13", "13"),
        ("　統整活動一", "篇章閱讀（讀信與回信）、審題立意、朗讀技巧", "4", "10", "10"),
        ("第二單元：天空的奇想", "探索天空奇想、科學求真、克服挫折與讀報素養", "24", "53", "53"),
        ("　第 4 課 飛行夢", "學習萊特兄弟面對挫折、堅持改良與實現夢想", "6", "13", "13"),
        ("　第 5 課 月光下", "賞析河邊賞月意境、掌握美景四字語詞與比喻", "6", "13", "13"),
        ("　第 6 課 又遠又近的月亮", "登月報導閱讀、掌握新聞六何要素與科普說明", "6", "14", "14"),
        ("　統整活動二", "六何法閱讀策略、寫作選材、讀報應用", "3", "7", "7"),
        ("　總複習與定期評量", "全冊前六課統整複習、情境素養整合評量", "3", "6", "6"),
    ]
    
    # Fill Table 1 rows
    for idx, row_data in enumerate(t1_data):
        r_idx = idx + 1
        row = t1.rows[r_idx]
        is_unit_header = "單元：" in row_data[0]
        set_cell_text(row.cells[0], row_data[0], bold=is_unit_header, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(row.cells[1], row_data[1], bold=is_unit_header, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(row.cells[2], row_data[2], bold=is_unit_header, size=10)
        set_cell_text(row.cells[3], row_data[3], bold=is_unit_header, size=10)
        set_cell_text(row.cells[4], row_data[4], bold=is_unit_header, size=10)
        
    # Clear any extra rows before Total (rows 12, 13)
    for r_idx in [12, 13]:
        row = t1.rows[r_idx]
        for c in row.cells:
            set_cell_text(c, "", size=9)
            
    # Row 14: 合計
    total_row = t1.rows[14]
    set_cell_text(total_row.cells[0], "合　　計", bold=True, size=10.5)
    # Note: total_row.cells[1] is merged with cell 0
    set_cell_text(total_row.cells[2], "45 節", bold=True, size=10.5)
    set_cell_text(total_row.cells[3], "100 ％", bold=True, size=10.5)
    set_cell_text(total_row.cells[4], "100 ％", bold=True, size=10.5)

    # 3. Table 2: 測驗的題型與配分
    t2 = doc.tables[1]
    t2_headers = ["題型項目", "一、國字注音", "二、改錯字", "三、選擇題", "四、成語填空", "五、句型練習", "六、閱讀理解"]
    t2_scores = ["配　　分", "28 分", "12 分", "20 分", "10 分", "14 分", "16 分"]
    for c_idx in range(7):
        set_cell_text(t2.rows[0].cells[c_idx], t2_headers[c_idx], bold=True, size=10)
        set_cell_text(t2.rows[1].cells[c_idx], t2_scores[c_idx], bold=True, size=10)

    # 4. Table 3: 試題分析雙向細目表
    t3 = doc.tables[2]
    # Header row is already set: 認知層次 單元名稱 | 記憶 | 理解 | 應用 | 分析 | 評鑑 | 創作 | 合計
    # Row 1: 第一單元
    # Row 2: 第二單元
    # Row 3: 清空或保留
    # Row 4: 清空或保留
    # Row 5: 合計
    t3_data = [
        ("第一單元：好山好水好故鄉\n（第1~3課、統整一）", "18 分\n12(1)+2(2)+1(2)", "12 分\n2(2)+2(4)", "12 分\n3(2)+1(2)+1(4)", "2 分\n1(2)", "0 分", "3 分\n1(3)", "47 分"),
        ("第二單元：天空的奇想\n（第4~6課、統整二）", "26 分\n16(1)+4(2)+1(2)", "6 分\n3(2)", "8 分\n2(2)+1(4)", "4 分\n1(4)", "6 分\n1(2)+1(4)", "3 分\n1(3)", "53 分"),
    ]
    
    for r_idx, (u_title, m_mem, m_und, m_app, m_ana, m_eva, m_cre, m_tot) in enumerate(t3_data, start=1):
        row = t3.rows[r_idx]
        set_cell_text(row.cells[0], u_title, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], m_mem, size=9)
        set_cell_text(row.cells[2], m_und, size=9)
        set_cell_text(row.cells[3], m_app, size=9)
        set_cell_text(row.cells[4], m_ana, size=9)
        set_cell_text(row.cells[5], m_eva, size=9)
        set_cell_text(row.cells[6], m_cre, size=9)
        set_cell_text(row.cells[7], m_tot, bold=True, size=10)

    # Rows 3 and 4 can be cleaned up or merged / removed.
    # In template, rows 3 and 4 are empty. Let's delete row 4 and row 3 so table is compact!
    # In python-docx, to delete a row: row._tr.getparent().remove(row._tr)
    t3._tbl.remove(t3.rows[4]._tr)
    t3._tbl.remove(t3.rows[3]._tr)
    
    # Now row 3 is 合計!
    total_t3 = t3.rows[3]
    set_cell_text(total_t3.cells[0], "合　　計", bold=True, size=10.5)
    set_cell_text(total_t3.cells[1], "44 分\n(44%)", bold=True, size=9.5)
    set_cell_text(total_t3.cells[2], "18 分\n(18%)", bold=True, size=9.5)
    set_cell_text(total_t3.cells[3], "20 分\n(20%)", bold=True, size=9.5)
    set_cell_text(total_t3.cells[4], "6 分\n(6%)", bold=True, size=9.5)
    set_cell_text(total_t3.cells[5], "6 分\n(6%)", bold=True, size=9.5)
    set_cell_text(total_t3.cells[6], "6 分\n(6%)", bold=True, size=9.5)
    set_cell_text(total_t3.cells[7], "100 分\n(100%)", bold=True, size=10.5)

    out_docx = r"d:\opencode\115-\filled_spec.docx"
    doc.save(out_docx)
    print("Saved filled_spec.docx successfully.")

fill_spec_table()
